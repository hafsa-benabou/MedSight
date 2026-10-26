"""Briques communes du pipeline d'import : provenance, lecture de fichiers, parsing, upserts."""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models as m
from app.core.ean import normalize_ean
from app.nlp.text import normalize_text

log = logging.getLogger("medisight.ingestion")

_AMMPS_RIGHTS = ("Consultation publique. Le site indique « tous droits réservés » : extraction massive ou republication "
                 "à faire autoriser par l'AMMPS. Échantillons saisis manuellement à des fins pédagogiques.")
SOURCES: dict[str, dict[str, Any]] = {
    "ammps_medicaments": dict(
        name="AMMPS — Base de données des médicaments", publisher="Agence Marocaine du Médicament et des Produits de Santé",
        url="https://www.ammps.gov.ma/recherche-medicaments", license="© AMMPS — tous droits réservés",
        reuse_conditions=_AMMPS_RIGHTS, is_official=True, default_data_status="official",
        notes="Base affichant une mise à jour au 01/10/2026 (9 936 médicaments). Référence prioritaire pour les prix."),
    "ammps_generiques": dict(
        name="AMMPS — Répertoire Marocain des Médicaments Génériques", publisher="AMMPS",
        url="https://www.ammps.gov.ma/repertoire-medicaments-generiques", license="© AMMPS — tous droits réservés",
        reuse_conditions=_AMMPS_RIGHTS, is_official=True, default_data_status="official",
        notes="Édition projet (janvier 2026), évolutive : vérifier la dernière version avant usage."),
    "ammps_pharmacies": dict(
        name="AMMPS — Base de données des pharmacies", publisher="AMMPS",
        url="https://www.ammps.gov.ma/bases-de-donnees/pharmacies", license="© AMMPS — tous droits réservés",
        reuse_conditions=_AMMPS_RIGHTS, is_official=True, default_data_status="official",
        notes="Nom et ville uniquement (pas d'adresse, téléphone ni coordonnées)."),
    "cnops_referentiel": dict(
        name="CNOPS — Référentiel des médicaments (data.gov.ma)", publisher="Caisse Nationale des Organismes de Prévoyance Sociale",
        url="https://data.gov.ma/data/dataset/referentiel-des-medicaments", license="ODbL 1.0",
        license_url="https://opendatacommons.org/licenses/odbl/1-0/",
        reuse_conditions="Réutilisation libre avec attribution ; partage à l'identique des bases dérivées (ODbL).",
        is_official=True, default_data_status="official",
        notes="Fichier ref-des-medicaments-cnops-2014.xlsx (ressource mise à jour le 18/12/2017) : prix potentiellement obsolètes."),
    "osm_pharmacies": dict(
        name="OpenStreetMap — pharmacies (amenity=pharmacy)", publisher="Contributeurs OpenStreetMap",
        url="https://www.openstreetmap.org/copyright", license="ODbL 1.0",
        license_url="https://opendatacommons.org/licenses/odbl/1-0/",
        reuse_conditions="Attribution obligatoire : « © les contributeurs d'OpenStreetMap ».",
        is_official=False, default_data_status="unverified"),
    "duty_manual": dict(
        name="Pharmacies de garde — import manuel", publisher="Équipe MEDISIGHT", license="Selon la source saisie",
        reuse_conditions="Vérifier les conditions de réutilisation de chaque liste de garde avant import.",
        is_official=False, default_data_status="unverified"),
    "rcp_documents": dict(
        name="RCP / notices importés", publisher="Titulaires d'AMM (via documents fournis)", license="Document source",
        reuse_conditions="Extraits de sections citées avec leur source ; usage d'information uniquement.",
        is_official=False, default_data_status="unverified"),
    "medbase_api": dict(
        name="MedBase API", publisher="MedBase", url="https://medbase.ma/api/documentation",
        license="Accès sur clé API (conditions MedBase)", is_official=False, default_data_status="reference"),
    "ddinter": dict(
        name="DDInter 2.0", publisher="SCBDD", url="https://ddinter2.scbdd.com", license="CC BY-NC-SA 4.0",
        license_url="https://creativecommons.org/licenses/by-nc-sa/4.0/",
        reuse_conditions="Usage non commercial, attribution, partage à l'identique.", is_official=False,
        default_data_status="reference"),
    "manual_interactions": dict(
        name="Interactions — saisie manuelle sourcée", publisher="Équipe MEDISIGHT", license="Selon la référence citée",
        is_official=False, default_data_status="unverified"),
    "demo_simulated": dict(
        name="Jeu de démonstration MEDISIGHT (simulé)", publisher="Équipe MEDISIGHT", license="Données fictives",
        reuse_conditions="Ne jamais présenter comme des données réelles.", is_official=False,
        default_data_status="simulated"),
    "team_nlp_dataset": dict(
        name="Corpus d'énoncés annotés MEDISIGHT", publisher="Équipe MEDISIGHT", license="Interne au projet",
        notes="Énoncés rédigés par l'équipe (synthétiques), à compléter par des énoncés réels anonymisés et consentis.",
        is_official=False, default_data_status="unverified"),
}


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def get_source(db: Session, code: str) -> m.DataSource:
    src = db.scalar(select(m.DataSource).where(m.DataSource.code == code))
    if src is None:
        if code not in SOURCES:
            raise ValueError(f"Source inconnue : {code}")
        src = m.DataSource(code=code, **SOURCES[code])
        db.add(src)
        db.flush()
    return src


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


@dataclass
class ImportContext:
    db: Session
    source: m.DataSource
    run: m.DataImport
    issues: list[dict[str, Any]] = field(default_factory=list)

    def warn(self, message: str, row: int | None = None) -> None:
        if len(self.issues) < 300:
            self.issues.append({"row": row, "level": "warning", "message": message})

    def reject(self, message: str, row: int | None = None) -> None:
        self.run.rows_rejected += 1
        if len(self.issues) < 300:
            self.issues.append({"row": row, "level": "rejected", "message": message})

    def summary(self) -> dict[str, Any]:
        r = self.run
        return {"import_id": r.id, "source": self.source.code, "status": r.status, "rows_read": r.rows_read,
                "inserted": r.rows_inserted, "updated": r.rows_updated, "rejected": r.rows_rejected,
                "warnings": sum(1 for i in self.issues if i["level"] == "warning")}


@contextmanager
def import_run(db: Session, source_code: str, *, file_path: str | Path | None = None, method: str = "file",
               version_date: date | None = None, triggered_by: str = "cli") -> Iterator[ImportContext]:
    source = get_source(db, source_code)
    run = m.DataImport(source_id=source.id, status="running", method=method,
                       file_name=Path(file_path).name if file_path else None,
                       file_sha256=file_sha256(file_path) if file_path else None,
                       source_version_date=version_date, started_at=utcnow(), rows_read=0, rows_inserted=0,
                       rows_updated=0, rows_rejected=0, report={}, triggered_by=triggered_by)
    db.add(run)
    db.flush()
    ctx = ImportContext(db=db, source=source, run=run)
    try:
        yield ctx
    except Exception as exc:
        db.rollback()
        src = get_source(db, source_code)
        db.add(m.DataImport(source_id=src.id, status="failed", method=method, file_name=run.file_name,
                            source_version_date=version_date, started_at=run.started_at, finished_at=utcnow(),
                            report={"error": str(exc)[:500]}, triggered_by=triggered_by))
        db.commit()
        raise
    else:
        run.finished_at = utcnow()
        run.status = "partial" if run.rows_rejected else "success"
        run.report = {"issues": ctx.issues, "issue_count": len(ctx.issues)}
        db.commit()


# ---------- Lecture de fichiers (CSV / XLSX) avec détection de la ligne d'en-têtes ----------

def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", normalize_text(value)).strip()


def _iter_raw_rows(path: Path, sheet: str | None = None) -> Iterator[list[Any]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        for row in ws.iter_rows(values_only=True):
            yield list(row)
        wb.close()
        return
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    yield from csv.reader(io.StringIO(text), delimiter=delimiter)


def read_rows(path: str | Path, expected: dict[str, list[str]], *, min_matches: int = 3,
              sheet: str | None = None) -> tuple[dict[str, int], list[tuple[int, list[Any]]]]:
    path = Path(path)
    raw = list(_iter_raw_rows(path, sheet))
    wanted = {f: {_norm_header(s) for s in syns} for f, syns in expected.items()}
    for i, row in enumerate(raw[:25]):
        headers = [_norm_header(c) for c in row]
        mapping: dict[str, int] = {}
        for fld, syns in wanted.items():
            for idx, h in enumerate(headers):
                if h and h in syns and idx not in mapping.values():
                    mapping[fld] = idx
                    break
        if len(mapping) >= min_matches:
            data = [(i + j + 2, r) for j, r in enumerate(raw[i + 1:]) if any(c not in (None, "") for c in r)]
            return mapping, data
    raise ValueError(f"Ligne d'en-têtes introuvable dans {path.name}. Colonnes attendues : {sorted(expected)}")


def cell(row: list[Any], mapping: dict[str, int], fld: str) -> Any:
    idx = mapping.get(fld)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


# ---------- Parsing ----------

def parse_price(raw: Any) -> tuple[Decimal | None, str]:
    """« 4.303.00 DH » -> 4303.00 ; « 10,70 » -> 10.70 ; « SANS PPV » -> (None, no_price)."""
    if raw is None:
        return None, "unknown"
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)).quantize(Decimal("0.01")), "priced"
    s = str(raw).strip().upper()
    if not s or s in {"-", "N/A", "NA", "ND"}:
        return None, "unknown"
    if "SANS" in s:
        return None, "no_price"
    s = re.sub(r"(DHS|DH|MAD|درهم)", "", s).replace(" ", "").replace("\u00a0", "").replace(",", ".")
    if s.count(".") > 1:
        head, _, tail = s.rpartition(".")
        s = head.replace(".", "") + "." + tail
    try:
        return Decimal(s).quantize(Decimal("0.01")), "priced"
    except InvalidOperation:
        return None, "unknown"


def parse_percent(raw: Any) -> Decimal | None:
    if raw is None or raw == "":
        return None
    try:
        value = Decimal(str(raw).replace("%", "").replace(",", ".").strip())
    except InvalidOperation:
        return None
    return value * 100 if value <= 1 else value


def parse_date(raw: Any) -> date | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def split_substances(raw: Any) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"\s*//\s*|\s*\+\s*", str(raw).strip())
    return [re.sub(r"\s+", " ", p).strip().upper() for p in parts if p and p.strip() and p.strip() != "-"]


@dataclass
class Strength:
    text: str | None
    value: Decimal | None = None
    unit: str | None = None
    mg: Decimal | None = None


_UNIT_MAP = {"mg": "mg", "mgr": "mg", "g": "g", "gr": "g", "µg": "µg", "μg": "µg", "mcg": "µg", "ug": "µg",
             "ml": "ml", "l": "l", "ui": "UI", "u": "U", "mui": "MUI", "mmol": "mmol", "meq": "mEq", "mbq": "MBq",
             "gbq": "GBq", "%": "%", "uceip": "UCEIP"}
_MASS = {"mg": Decimal(1), "g": Decimal(1000), "µg": Decimal("0.001")}
_STRENGTH_RE = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*([a-zA-Zµμ%]+)?\s*(?:/\s*(.+))?$")


def parse_strength(text: str | None) -> Strength:
    t = (text or "").strip()
    match = _STRENGTH_RE.match(t)
    if not match:
        return Strength(t or None)
    value = Decimal(match.group(1).replace(",", "."))
    unit_raw = (match.group(2) or "").lower()
    unit = _UNIT_MAP.get(unit_raw, unit_raw or None)
    if match.group(3):
        unit = f"{unit or ''}/{match.group(3).strip().lower()}"
    mg = value * _MASS[unit] if unit in _MASS else None
    return Strength(t, value, unit, mg)


def split_strengths(dosage: str | None, n: int) -> list[Strength | None]:
    """« 500 MG / 50 MG » avec 2 substances -> 2 dosages ; « 80 MG / ML » avec 1 substance -> concentration."""
    if not dosage:
        return []
    if n <= 1:
        return [parse_strength(dosage)]
    parts = [p.strip() for p in re.split(r"\s*/\s*", dosage) if p.strip()]
    if len(parts) == n:
        return [parse_strength(p) for p in parts]
    return [Strength(dosage)] + [None] * (n - 1)


_FORM_RULES = [
    (r"effervescent", "comprime_effervescent"), (r"orodispersible", "comprime_orodispersible"),
    (r"\bcompr|\bcompeime|\bcomrime", "comprime"), (r"\bgelule", "gelule"), (r"\bcapsule", "capsule"),
    (r"\bsuppositoire", "suppositoire"), (r"\bcollyre|ophtalmique|oculaire", "collyre"), (r"\bsirop", "sirop"),
    (r"suspension (buvable|orale)", "suspension_buvable"), (r"solution buvable|gouttes? buvables?", "solution_buvable"),
    (r"injectable|injection|perfusion|lyophilisat|lyophylisat|parenteral", "injectable"), (r"\bsachet", "sachet"),
    (r"\bgranule", "granules"), (r"\bpommade", "pommade"), (r"\bcreme", "creme"), (r"\bgel\b", "gel"),
    (r"inhal|nebuli", "inhalation"), (r"\bspray|pulverisation", "spray"), (r"\bovule", "ovule"),
    (r"\bpatch|transdermique", "patch"), (r"\bgoutte", "gouttes"), (r"\bpoudre", "poudre"),
    (r"\bsolution", "solution"), (r"\blotion", "lotion"), (r"shampo", "shampooing"),
]


def canonical_form(label: str | None) -> str | None:
    norm = normalize_text(label or "")
    if not norm:
        return None
    for pattern, canonical in _FORM_RULES:
        if re.search(pattern, norm):
            return canonical
    return "autre"


def marketed_flag(status: str | None) -> bool | None:
    n = normalize_text(status or "")
    if not n:
        return None
    if any(k in n for k in ("non commercialis", "sans prix", "retire", "suspendu", "arret")):
        return False
    return True if "commercialis" in n else None


def norm_princeps(raw: Any) -> str | None:
    n = normalize_text(raw or "").replace(" ", "")
    return {"pp": "P", "p": "P", "princeps": "P", "gn": "G", "g": "G", "generique": "G", "i": "I"}.get(n)


def medicine_key(brand: str, dosage: str | None, form: str | None) -> str:
    return "|".join(normalize_text(x or "") for x in (brand, dosage, form))


# ---------- Upserts ----------

class RefCache:
    def __init__(self, db: Session):
        self.db, self._forms, self._labs, self._classes, self._ings = db, {}, {}, {}, {}

    def form(self, label: str | None) -> m.PharmaceuticalForm | None:
        if not label:
            return None
        key = re.sub(r"\s+", " ", label).strip().upper()
        if key not in self._forms:
            obj = self.db.scalar(select(m.PharmaceuticalForm).where(m.PharmaceuticalForm.label == key))
            if obj is None:
                obj = m.PharmaceuticalForm(label=key, canonical=canonical_form(key))
                self.db.add(obj)
                self.db.flush()
            self._forms[key] = obj
        return self._forms[key]

    def lab(self, name: str | None) -> m.Laboratory | None:
        if not name:
            return None
        key = re.sub(r"\s+", " ", name).strip().upper()
        if key not in self._labs:
            obj = self.db.scalar(select(m.Laboratory).where(m.Laboratory.name == key))
            if obj is None:
                obj = m.Laboratory(name=key, name_normalized=normalize_text(key))
                self.db.add(obj)
                self.db.flush()
            self._labs[key] = obj
        return self._labs[key]

    def therapeutic_class(self, name: str | None) -> m.TherapeuticClass | None:
        if not name:
            return None
        key = re.sub(r"\s+", " ", name).strip().upper()
        if key not in self._classes:
            obj = self.db.scalar(select(m.TherapeuticClass).where(m.TherapeuticClass.name == key))
            if obj is None:
                obj = m.TherapeuticClass(name=key)
                self.db.add(obj)
                self.db.flush()
            self._classes[key] = obj
        return self._classes[key]

    def _alias_map(self) -> dict[str, int]:
        if not hasattr(self, "_aliases"):
            self._aliases = {}
            for ing in self.db.scalars(select(m.ActiveIngredient)):
                for syn in ing.synonyms or []:
                    self._aliases.setdefault(normalize_text(syn), ing.id)
        return self._aliases

    def ingredient(self, name: str) -> m.ActiveIngredient:
        key = normalize_text(name)
        if key not in self._ings:
            obj = self.db.scalar(select(m.ActiveIngredient).where(m.ActiveIngredient.name_normalized == key))
            if obj is None and key in self._alias_map():  # ex. « VITAMINE B1 (THIAMINE) » -> THIAMINE
                obj = self.db.get(m.ActiveIngredient, self._alias_map()[key])
            if obj is None:
                obj = m.ActiveIngredient(name=re.sub(r"\s+", " ", name).strip().upper(), name_normalized=key, synonyms=[])
                self.db.add(obj)
                self.db.flush()
            self._ings[key] = obj
        return self._ings[key]


def upsert_medicine(ctx: ImportContext, cache: RefCache, *, brand: str, dosage: str | None, form_label: str | None,
                    substances: list[str], laboratory: str | None = None, therapeutic_class: str | None = None,
                    princeps: str | None = None, amm_status: str | None = None, marketing_status: str | None = None,
                    rcp_url: str | None = None, atc: str | None = None,
                    data_status: str | None = None) -> tuple[m.Medicine, bool]:
    key = medicine_key(brand, dosage, form_label)
    status = data_status or ctx.source.default_data_status
    med = ctx.db.scalar(select(m.Medicine).where(m.Medicine.source_record_key == key))
    created = med is None
    if created:
        med = m.Medicine(source_record_key=key, brand_name=brand.strip(), brand_normalized=normalize_text(brand),
                         data_status=status, source_id=ctx.source.id)
        ctx.db.add(med)
    elif m.STATUS_RANK[status] < m.STATUS_RANK[med.data_status]:
        return med, False  # ne jamais dégrader une donnée officielle par une donnée simulée/non vérifiée
    med.dosage_text = dosage
    med.form = cache.form(form_label) or med.form
    med.laboratory = cache.lab(laboratory) or med.laboratory
    med.therapeutic_class = cache.therapeutic_class(therapeutic_class) or med.therapeutic_class
    med.princeps_status = princeps or med.princeps_status
    med.amm_status = amm_status or med.amm_status
    if marketing_status:
        med.marketing_status, med.is_marketed = marketing_status, marketed_flag(marketing_status)
    med.rcp_url = rcp_url or med.rcp_url
    med.atc_code = atc or med.atc_code
    med.data_status, med.source_id = status, ctx.source.id
    if substances:
        strengths = split_strengths(dosage, len(substances))
        new = [(cache.ingredient(name), strengths[i] if i < len(strengths) else None) for i, name in enumerate(substances)]
        old = [(mi.ingredient_id, mi.strength_text) for mi in med.ingredients]
        if created or old != [(ing.id, st.text if st else None) for ing, st in new]:
            med.ingredients.clear()
            ctx.db.flush()
            for pos, (ing, st) in enumerate(new):
                med.ingredients.append(m.MedicineIngredient(
                    position=pos, ingredient=ing, strength_text=st.text if st else None,
                    strength_value=st.value if st else None, strength_unit=st.unit if st else None,
                    strength_mg=st.mg if st else None))
    med.search_text = normalize_text(" ".join([brand, dosage or "", " ".join(substances), form_label or ""]))
    ctx.db.flush()
    return med, created


def upsert_presentation(ctx: ImportContext, med: m.Medicine, *, label: str | None, ppv: Decimal | None = None,
                        ph: Decimal | None = None, pfht: Decimal | None = None, price_status: str = "unknown",
                        price_as_of: date | None = None, ean: Any = None,
                        data_status: str | None = None) -> tuple[m.MedicinePresentation, bool]:
    text = (label or "PRÉSENTATION NON PRÉCISÉE").strip()
    key = normalize_text(text)
    status = data_status or ctx.source.default_data_status
    pres = next((p for p in med.presentations if p.label_normalized == key), None)
    created = pres is None
    if created:
        pres = m.MedicinePresentation(label=text, label_normalized=key, data_status=status, source_id=ctx.source.id,
                                      price_status="unknown")
        med.presentations.append(pres)
    if ean is not None:
        code, est = normalize_ean(ean)
        if code and est != "valid":
            ctx.warn(f"Code {code} ({est}) ignoré pour {med.brand_name} {text}")
        elif code and pres.ean13 != code:
            pres.ean13, pres.ean_status = code, est
    if price_status != "unknown":
        allowed = created or m.STATUS_RANK[status] >= m.STATUS_RANK[pres.data_status]
        newer = pres.price_as_of is None or price_as_of is None or price_as_of >= pres.price_as_of
        if allowed and (newer or pres.price_status == "unknown"):
            if ppv is not None and pres.ppv_mad is not None and ppv != pres.ppv_mad:
                ctx.warn(f"Prix {med.brand_name} {text} : {pres.ppv_mad} remplacé par {ppv} (source plus récente)")
            pres.ppv_mad, pres.ph_mad, pres.pfht_mad = ppv, ph, pfht
            pres.price_status, pres.price_as_of = price_status, price_as_of
            pres.data_status, pres.source_id = status, ctx.source.id
        elif ppv is not None and pres.ppv_mad is not None and ppv != pres.ppv_mad:
            ctx.warn(f"Prix {med.brand_name} {text} : valeur {ppv} plus ancienne ignorée (actuelle {pres.ppv_mad})")
    ctx.db.flush()
    return pres, created
